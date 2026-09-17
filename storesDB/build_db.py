"""Build an enriched SQLite database from stores.xlsx + websites.xlsx.

Each store row from stores.xlsx is enriched with website data from
websites.xlsx by matching the website URL as a prefix of the store_url
(e.g. https://cancun.shopdutyfree.com/ matches
https://cancun.shopdutyfree.com/en/13/).
"""
import os
import sqlite3
import openpyxl

BASE_DIR = os.path.dirname(__file__)
STORES_XLSX = os.path.join(BASE_DIR, '..', 'utils', 'stores.xlsx')
WEBSITES_XLSX = os.path.join(BASE_DIR, '..', 'utils', 'websites.xlsx')
DB_PATH = os.path.join(BASE_DIR, 'stores.db')

NULL_VALUES = {None, "", "NULL", "null"}


def clean(val):
    if val in NULL_VALUES:
        return None
    return str(val).replace("\xa0", " ").strip()


def norm(url):
    if url in NULL_VALUES:
        return ""
    return str(url).strip().rstrip("/").lower()


def read_sheet_rows(ws):
    rows = ws.iter_rows(values_only=True)
    try:
        headers = [str(h).strip() if h is not None else "" for h in next(rows)]
    except StopIteration:
        return []
    return [dict(zip(headers, row)) for row in rows]


def read_stores():
    wb = openpyxl.load_workbook(STORES_XLSX, read_only=True, data_only=True)
    data = read_sheet_rows(wb.active)
    wb.close()
    return data


def map_website_row(row):
    """Map a websites.xlsx row from either known schema to a common shape."""
    # Standard schema (most divisions).
    if "URL" in row:
        url = row.get("URL")
        return {
            "website_name": clean(row.get("website_name")),
            "website_code": clean(row.get("website_code")),
            "website_url": clean(url),
            "has_active_store": clean(row.get("has_active_store")),
            "iata_code": clean(row.get("Código IATA")),
            "country": clean(row.get("País")),
            "region": clean(row.get("Región")),
            "key": norm(url),
        }
    # Alternate schema (DIV7, DIV10) has no status column; assume active.
    if "Website Url" in row:
        url = row.get("Website Url")
        return {
            "website_name": clean(row.get("Store Name")),
            "website_code": clean(row.get("Store Code")),
            "website_url": clean(url),
            "has_active_store": "Active",
            "iata_code": clean(row.get("Código IATA")),
            "country": clean(row.get("País")),
            "region": clean(row.get("Región")),
            "key": norm(url),
        }
    return None


def load_websites():
    wb = openpyxl.load_workbook(WEBSITES_XLSX, read_only=True, data_only=True)
    websites = []
    for name in wb.sheetnames:
        for row in read_sheet_rows(wb[name]):
            w = map_website_row(row)
            if not w or not w["key"]:
                continue
            websites.append(w)
    wb.close()
    # Longest key first so the most specific URL wins on prefix match.
    websites.sort(key=lambda x: len(x["key"]), reverse=True)
    return websites


def match_website(store_url, websites):
    su = norm(store_url)
    if not su:
        return None
    for w in websites:
        if su.startswith(w["key"]):
            return w
    return None


def build():
    websites = load_websites()
    stores = read_stores()

    if os.path.exists(DB_PATH):
        os.remove(DB_PATH)
    conn = sqlite3.connect(DB_PATH)
    conn.execute("""
        CREATE TABLE stores (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            division TEXT,
            store_id INTEGER,
            code TEXT,
            name TEXT,
            website_id INTEGER,
            store_url TEXT,
            website_name TEXT,
            website_code TEXT,
            website_url TEXT,
            has_active_store TEXT,
            iata_code TEXT,
            country TEXT,
            region TEXT
        )
    """)

    matched = 0
    for s in stores:
        w = match_website(s.get("store_url"), websites) or {}
        if w:
            matched += 1
        conn.execute("""
            INSERT INTO stores (
                division, store_id, code, name, website_id, store_url,
                website_name, website_code, website_url,
                has_active_store, iata_code, country, region
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            s.get("division"), s.get("store_id"), s.get("code"), s.get("name"),
            s.get("website_id"), s.get("store_url"),
            w.get("website_name"), w.get("website_code"), w.get("website_url"),
            w.get("has_active_store"), w.get("iata_code"),
            w.get("country"), w.get("region"),
        ))

    conn.commit()
    conn.close()
    print(f"DB creada en {DB_PATH}")
    print(f"  {len(stores)} stores, {matched} enriquecidas con datos de website")


if __name__ == "__main__":
    build()
